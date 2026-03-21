#!/usr/bin/env python3
"""Generate the Embassy Consular Portal demo pitch deck."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand colors
GREEN = RGBColor(0x0A, 0x5C, 0x36)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


def add_blank_slide():
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def add_green_bar(slide, top=0, height=Inches(0.08)):
    """Add a thin green accent bar."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, top, W, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREEN
    shape.line.fill.background()


def add_bottom_bar(slide):
    """Subtle bottom accent line."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, H - Inches(0.04), W, Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREEN
    shape.line.fill.background()


def set_text(tf, text, size=18, bold=False, color=DARK, alignment=PP_ALIGN.LEFT):
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return p


def add_title_text(slide, text, left=Inches(0.9), top=Inches(0.5), width=Inches(11), size=32):
    txBox = slide.shapes.add_textbox(left, top, width, Inches(0.7))
    tf = txBox.text_frame
    set_text(tf, text, size=size, bold=True, color=GREEN)
    return txBox


def add_subtitle_line(slide, top=Inches(1.15)):
    """Green line under title."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.9), top, Inches(2.5), Inches(0.04)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GREEN
    shape.line.fill.background()


def add_bullet_slide(slide, bullets, start_top=Inches(1.6), left=Inches(1.2), width=Inches(10)):
    for i, bullet in enumerate(bullets):
        top = start_top + Inches(i * 0.7)
        # Bullet dot
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, left, top + Inches(0.08), Inches(0.12), Inches(0.12)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = GREEN
        dot.line.fill.background()
        # Text
        txBox = slide.shapes.add_textbox(left + Inches(0.35), top - Inches(0.02), width, Inches(0.5))
        tf = txBox.text_frame
        set_text(tf, bullet, size=20, color=DARK)


# =====================================================
# SLIDE 1 — TITLE
# =====================================================
slide = add_blank_slide()
add_green_bar(slide, top=0, height=Inches(0.1))
add_bottom_bar(slide)

# Green left accent block
shape = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, Inches(2.0), Inches(0.15), Inches(3.5)
)
shape.fill.solid()
shape.fill.fore_color.rgb = GREEN
shape.line.fill.background()

# Main title
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10), Inches(1.2))
tf = txBox.text_frame
tf.word_wrap = True
set_text(tf, "Consular Workflow Improvement", size=40, bold=True, color=GREEN)

# Subtitle
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(3.3), Inches(10), Inches(0.6))
tf = txBox.text_frame
set_text(tf, "Pilot Proposal", size=28, bold=False, color=DARK)

# Embassy line
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(4.1), Inches(10), Inches(0.5))
tf = txBox.text_frame
set_text(tf, "Embassy of Nigeria, Paris", size=20, color=GRAY)

# Author placeholder
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(5.5), Inches(6), Inches(0.8))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "[Your Name]\n[Date]"
run.font.size = Pt(16)
run.font.color.rgb = GRAY
run.font.name = "Calibri"


# =====================================================
# SLIDE 2 — THE PROBLEM
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "Current Challenges")
add_subtitle_line(slide)

bullets = [
    "High volume of passport applications processed manually",
    "Frequent applicant follow-ups via email and phone calls",
    "Limited visibility into application status for applicants",
    "Manual document handling increases error risk",
    "Delays caused by incomplete or incorrect submissions",
]
add_bullet_slide(slide, bullets)


# =====================================================
# SLIDE 3 — IMPACT
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "Operational Impact")
add_subtitle_line(slide)

bullets = [
    "Increased workload for consular staff due to repetitive inquiries",
    "Significant time spent on manual communication with applicants",
    "Difficulty tracking application progress across multiple cases",
    "Reduced overall efficiency in document processing workflow",
]
add_bullet_slide(slide, bullets)


# =====================================================
# SLIDE 4 — PROPOSED SOLUTION
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "Proposed Solution (Pilot)")
add_subtitle_line(slide)

bullets = [
    "Simple digital portal for applicants and consular staff",
    "Structured document submission with completeness checks",
    "Real-time application status tracking for applicants",
    "Centralized review dashboard for consular officers",
]
add_bullet_slide(slide, bullets)


# =====================================================
# SLIDE 5 — HOW IT WORKS
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "How the System Works")
add_subtitle_line(slide)

steps = [
    ("1", "Applicant submits\napplication + documents"),
    ("2", "Embassy reviews\ndocuments"),
    ("3", "Status is updated\nin real-time"),
    ("4", "Applicant tracks\nprogress online"),
]

box_width = Inches(2.4)
box_height = Inches(2.2)
start_left = Inches(1.0)
arrow_width = Inches(0.6)
top = Inches(2.5)

for i, (num, text) in enumerate(steps):
    left = start_left + i * (box_width + arrow_width)

    # Box
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, box_width, box_height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = GREEN
    box.line.width = Pt(1.5)

    # Number circle
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, left + Inches(0.85), top + Inches(0.25), Inches(0.7), Inches(0.7)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = GREEN
    circle.line.fill.background()
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = num
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Step text
    txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(1.15), box_width - Inches(0.3), Inches(0.9))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(15)
    run.font.color.rgb = DARK
    run.font.name = "Calibri"

    # Arrow between boxes
    if i < 3:
        arrow_left = left + box_width + Inches(0.05)
        arrow_top = top + Inches(0.95)
        arrow = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_ARROW, arrow_left, arrow_top, Inches(0.5), Inches(0.3)
        )
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = GREEN
        arrow.line.fill.background()


# =====================================================
# SLIDE 6 — PROTOTYPE OVERVIEW (Screenshots placeholder)
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "Prototype Overview")
add_subtitle_line(slide)

# Three placeholder boxes for screenshots
labels = ["Application Portal", "Document Upload", "Admin Dashboard"]
for i, label in enumerate(labels):
    left = Inches(0.9) + i * Inches(4.1)
    top = Inches(1.8)
    w = Inches(3.8)
    h = Inches(4.5)

    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_BG
    box.line.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    box.line.width = Pt(1)

    # Placeholder text
    txBox = slide.shapes.add_textbox(left, top + Inches(1.8), w, Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"[Insert Screenshot]\n{label}"
    run.font.size = Pt(14)
    run.font.color.rgb = GRAY
    run.font.name = "Calibri"


# =====================================================
# SLIDE 7 — EXPECTED BENEFITS
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "Expected Benefits")
add_subtitle_line(slide)

bullets = [
    "Reduced applicant inquiries through self-service status tracking",
    "Improved processing efficiency with structured document submission",
    "Better document organization and centralized review workflow",
    "Increased transparency and trust for applicants",
]
add_bullet_slide(slide, bullets)


# =====================================================
# SLIDE 8 — PILOT PROPOSAL
# =====================================================
slide = add_blank_slide()
add_green_bar(slide)
add_bottom_bar(slide)
add_title_text(slide, "Pilot Proposal (Paris Embassy)")
add_subtitle_line(slide)

bullets = [
    "Limited scope: Paris consular section only",
    "No changes to existing embassy systems or infrastructure",
    "Low-risk implementation alongside current processes",
    "Short evaluation period: 4 to 8 weeks",
]
add_bullet_slide(slide, bullets)

# Additional note
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(10), Inches(0.5))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "This pilot is designed to complement, not replace, existing workflows."
run.font.size = Pt(16)
run.font.italic = True
run.font.color.rgb = GRAY
run.font.name = "Calibri"


# =====================================================
# SLIDE 9 — NEXT STEP
# =====================================================
slide = add_blank_slide()
add_green_bar(slide, top=0, height=Inches(0.1))
add_bottom_bar(slide)

# Green left accent block
shape = slide.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, Inches(2.5), Inches(0.15), Inches(2.5)
)
shape.fill.solid()
shape.fill.fore_color.rgb = GREEN
shape.line.fill.background()

# Title
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(2.7), Inches(10), Inches(0.8))
tf = txBox.text_frame
set_text(tf, "Next Step", size=36, bold=True, color=GREEN)

# Content
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(3.7), Inches(10), Inches(0.6))
tf = txBox.text_frame
set_text(tf, "Short demonstration and feedback session", size=24, color=DARK)

# Subtext
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(4.6), Inches(10), Inches(0.8))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "I am available at your convenience for a brief presentation\nof the prototype and to gather your feedback."
run.font.size = Pt(17)
run.font.color.rgb = GRAY
run.font.name = "Calibri"

# Contact placeholder
txBox = slide.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(6), Inches(0.8))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
run = p.add_run()
run.text = "[Your Name]  |  [Phone Number]  |  [Email]"
run.font.size = Pt(15)
run.font.color.rgb = GRAY
run.font.name = "Calibri"


# =====================================================
# SAVE
# =====================================================
output = "/home/user/nigeria-embassy-paris-portal/presentation/demo-pitch.pptx"
prs.save(output)
print(f"Presentation saved to: {output}")
print(f"Total slides: {len(prs.slides)}")
